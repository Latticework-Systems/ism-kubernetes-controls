#!/usr/bin/env python3
"""Join a provider's Cloud Security Controls Matrix (CSCM) to the coverage census.

The CSCM is the provider's confidential per-control responsibility allocation,
downloaded by the customer (for AWS, from AWS Artifact). This script reads the
customer's own copy and writes a private CSV. Never commit the input or output.

With --catalog and --classification, the output covers every numbered ISM
control that applies at that classification, not only the census, and triages
each one: inherited from the provider, marked not applicable by the provider,
the customer's, or absent from the provider matrix.

Requires openpyxl: python3 -m pip install openpyxl
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path
import re
import warnings

import yaml


ROOT = Path(__file__).resolve().parents[1]
COVERAGE = ROOT / "mapping/coverage.yaml"
PROVENANCE_LOCK = ROOT / "mapping/provenance.lock.yaml"
CLASSIFICATIONS = ["NC", "OS", "P", "S", "TS"]
COLUMNS = {
    "Identifier": "id",
    "Responsible Entity": "responsible_entity",
    "Consumer Control Type": "consumer_control_type",
    "Implementation Status": "implementation_status",
    "Consumer Guidance": "consumer_guidance",
}


def read_cscm(path: Path, sheet: str) -> dict[str, dict[str, str]]:
    try:
        import openpyxl
    except ImportError as error:
        raise SystemExit("import_cscm.py requires openpyxl: python3 -m pip install openpyxl") from error

    # Provider workbooks use Excel data validation, which openpyxl drops with a warning; only values are read.
    warnings.filterwarnings("ignore", message="Data Validation extension is not supported", module="openpyxl")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in workbook.sheetnames:
        raise SystemExit(f"{path} has no sheet {sheet!r}; found {workbook.sheetnames}")
    rows = workbook[sheet].iter_rows(values_only=True)
    header = next((row for row in rows if row and row[0] == "Identifier"), None)
    if not header:
        raise SystemExit(f"{path}:{sheet} has no header row starting with 'Identifier'")
    index = {COLUMNS[name]: header.index(name) for name in COLUMNS if name in header}
    missing = set(COLUMNS.values()) - set(index)
    if missing:
        raise SystemExit(f"{path}:{sheet} is missing columns: {sorted(missing)}")

    controls = {}
    for row in rows:
        raw = str(row[index["id"]] or "").strip()
        if re.fullmatch(r"[0-9]{1,4}", raw):
            controls[f"ISM-{raw.zfill(4)}"] = {
                field: str(row[column] or "").strip() for field, column in index.items() if field != "id"
            }
    return controls


def read_catalog(path: Path, classification: str) -> dict[str, dict[str, str]]:
    """Return numbered ISM controls that apply at a classification, with their headings."""
    expected = yaml.safe_load(PROVENANCE_LOCK.read_text())["sources"]["asd_ism_catalog"]["sha256"]
    actual = hashlib.sha256(path.read_bytes()).hexdigest()
    if actual != expected:
        raise SystemExit(f"{path} does not match the pinned ASD catalogue: want {expected}, got {actual}")

    controls: dict[str, dict[str, str]] = {}

    def walk(node: dict, headings: list[str]) -> None:
        for control in node.get("controls", []):
            ism_id = control["id"].upper()
            applicability = {prop["value"] for prop in control.get("props", []) if prop["name"] == "applicability"}
            if re.fullmatch(r"ISM-[0-9]{4}", ism_id) and classification in applicability:
                statement = next(part["prose"] for part in control["parts"] if part["name"] == "statement")
                controls[ism_id] = {
                    "guideline": headings[0] if headings else "",
                    "section": headings[1] if len(headings) > 1 else "",
                    "title": statement,
                }
        for group in node.get("groups", []):
            walk(group, headings + [group.get("title", "")])

    walk(json.loads(path.read_text())["catalog"], [])
    return controls


def triage(provider: dict | None) -> str:
    if not provider:
        return "not-in-provider-matrix"
    if provider["consumer_control_type"] == "Inherited":
        return "inherited"
    if "Not Applicable" in {provider["responsible_entity"], provider["consumer_control_type"]}:
        return "provider-not-applicable"
    return "customer"


def flag(control: dict, provider: dict | None, platform: str) -> str:
    if not provider:
        return "absent from provider matrix; check ISM version"
    inherited = provider["consumer_control_type"] == "Inherited"
    customer_layers = set(control["layers"]) - {"provider"}
    if platform == "eks-fargate":
        customer_layers -= {"node"}
    if inherited and customer_layers and control["coverage"] in {"automated", "buildable"}:
        return "provider marks inherited; confirm whether the customer layers are in its scope"
    if control["coverage"] == "provider-report" and not inherited:
        return "expected inherited; provider allocates responsibility to the customer"
    return ""


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("cscm", type=Path, help="customer-supplied CSCM workbook (.xlsx)")
    parser.add_argument("--out", type=Path, required=True, help="private CSV output, outside the repository or under .private/")
    parser.add_argument("--sheet", default="CCM")
    parser.add_argument("--platform", choices=["eks-ec2", "eks-fargate"], default="eks-ec2")
    parser.add_argument("--coverage", type=Path, default=COVERAGE)
    parser.add_argument("--catalog", type=Path, help="pinned ASD ISM catalogue JSON; enables whole-system triage")
    parser.add_argument("--classification", choices=CLASSIFICATIONS, help="system classification for --catalog")
    args = parser.parse_args()
    if bool(args.catalog) != bool(args.classification):
        raise SystemExit("--catalog and --classification must be used together")

    out = args.out.resolve()
    if out.is_relative_to(ROOT) and not out.is_relative_to(ROOT / ".private"):
        raise SystemExit("write the joined matrix outside the repository, or under .private/, so it is never committed")

    provider = read_cscm(args.cscm, args.sheet)
    census = {control["ism_id"]: control for control in yaml.safe_load(args.coverage.read_text())["controls"]}
    if args.catalog:
        catalog = read_catalog(args.catalog, args.classification)
        ids = sorted(catalog)
        excluded = sorted(set(census) - set(catalog))
    else:
        catalog = {}
        ids = sorted(census)
        excluded = []

    args.out.parent.mkdir(parents=True, exist_ok=True)
    triaged: dict[str, int] = {}
    reviewed: dict[str, int] = {}
    flagged = 0
    with args.out.open("w", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            [
                "ism_id",
                "guideline",
                "section",
                "title",
                "triage",
                "coverage",
                "layers",
                "collectors",
                "provider_responsible_entity",
                "provider_consumer_control_type",
                "provider_implementation_status",
                "provider_consumer_guidance",
                "review_flag",
            ]
        )
        for ism_id in ids:
            row = provider.get(ism_id)
            control = census.get(ism_id)
            bucket = triage(row)
            triaged[bucket] = triaged.get(bucket, 0) + 1
            note = flag(control, row, args.platform) if control else ""
            flagged += bool(note)
            if bucket == "customer":
                state = control["coverage"] if control else "unreviewed"
                reviewed[state] = reviewed.get(state, 0) + 1
            layers = [
                "provider" if args.platform == "eks-fargate" and layer == "node" else layer
                for layer in (control or {}).get("layers", [])
            ]
            heading = catalog.get(ism_id, {})
            title = heading.get("title") or control["title"]
            writer.writerow(
                [
                    ism_id,
                    heading.get("guideline", ""),
                    heading.get("section", ""),
                    title.split("\n")[0],
                    bucket,
                    control["coverage"] if control else "unreviewed",
                    " ".join(dict.fromkeys(layers)),
                    " ".join((control or {}).get("collectors", [])),
                    (row or {}).get("responsible_entity", ""),
                    (row or {}).get("consumer_control_type", ""),
                    (row or {}).get("implementation_status", ""),
                    (row or {}).get("consumer_guidance", ""),
                    note,
                ]
            )

    print(f"wrote {len(ids)} controls to {args.out}; {flagged} flagged for review")
    print("triage: " + ", ".join(f"{key} {value}" for key, value in sorted(triaged.items())))
    print("customer controls by coverage: " + ", ".join(f"{key} {value}" for key, value in sorted(reviewed.items())))
    if excluded:
        print(f"census controls not applicable at {args.classification}: {', '.join(excluded)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
